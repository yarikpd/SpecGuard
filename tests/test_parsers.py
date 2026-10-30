import io
from pathlib import Path
import docx
import pytest
from pypdf import PdfReader, PdfWriter

from specguard.rag.parsers import (
    DocxParser,
    DocumentParser,
    PdfParser,
    TextParser,
    get_parser,
    parse_document,
)
from specguard.rag.parser import SpecParser


def _build_minimal_pdf_bytes(text: str = "SpecGuard PDF Content") -> bytes:
    escaped_text = text.encode("ascii", errors="replace")
    stream = b"BT /F1 12 Tf 50 250 Td (" + escaped_text + b") Tj ET\n"
    stream_len = len(stream)
    return (
        b"%PDF-1.4\n"
        b"1 0 obj<</Type/Catalog/Pages 2 0 R>>endobj\n"
        b"2 0 obj<</Type/Pages/Kids[3 0 R]/Count 1>>endobj\n"
        b"3 0 obj<</Type/Page/Parent 2 0 R/MediaBox[0 0 300 300]/Contents 4 0 R/Resources<</Font<</F1 5 0 R>>>>>>endobj\n"
        b"4 0 obj<</Length " + str(stream_len).encode("ascii") + b">>stream\n"
        + stream
        + b"endstream\nendobj\n"
        b"5 0 obj<</Type/Font/Subtype/Type1/BaseFont/Helvetica>>endobj\n"
        b"xref\n0 6\n"
        b"0000000000 65535 f \n"
        b"0000000009 00000 n \n"
        b"0000000052 00000 n \n"
        b"0000000101 00000 n \n"
        b"0000000212 00000 n \n"
        b"0000000306 00000 n \n"
        b"trailer<</Size 6/Root 1 0 R>>\nstartxref\n377\n%%EOF"
    )


def test_document_parser_protocol_conformance():
    assert isinstance(TextParser(), DocumentParser)
    assert isinstance(DocxParser(), DocumentParser)
    assert isinstance(PdfParser(), DocumentParser)


def test_parse_text_utf8(tmp_path: Path):
    md_file = tmp_path / "spec.md"
    content = "# ТЗ SpecGuard\n- FR-001: Аудит кода\n- FR-002: Поддержка форматов"
    md_file.write_text(content, encoding="utf-8")

    text, warnings = parse_document(md_file)
    assert text == content
    assert warnings == []


def test_parse_text_cp1251_fallback(tmp_path: Path):
    txt_file = tmp_path / "legacy_spec.txt"
    content = "Требование 1: Обработка кодировки Windows-1251"
    txt_file.write_bytes(content.encode("cp1251"))

    text, warnings = parse_document(txt_file)
    assert text == content
    assert any("cp1251" in w for w in warnings)


def test_parse_empty_text_file(tmp_path: Path):
    empty_file = tmp_path / "empty.txt"
    empty_file.write_text("   \n  ", encoding="utf-8")

    text, warnings = parse_document(empty_file)
    assert text == "   \n  "
    assert any("не содержит значащего текста" in w for w in warnings)


def test_parse_docx_paragraphs_and_tables(tmp_path: Path):
    docx_file = tmp_path / "spec.docx"
    doc = docx.Document()
    doc.add_heading("Спецификация", level=1)
    doc.add_paragraph("Первый параграф требований.")

    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Идентификатор"
    table.cell(0, 1).text = "Описание"
    table.cell(1, 0).text = "REQ-101"
    table.cell(1, 1).text = "Интеграция с GitHub"

    doc.save(docx_file)

    text, warnings = parse_document(docx_file)
    assert warnings == []
    assert "Спецификация" in text
    assert "Первый параграф требований." in text
    assert "Идентификатор | Описание" in text
    assert "REQ-101 | Интеграция с GitHub" in text


def test_parse_docx_corrupted(tmp_path: Path):
    corrupt_docx = tmp_path / "corrupt.docx"
    corrupt_docx.write_bytes(b"This is completely invalid zip data")

    text, warnings = parse_document(corrupt_docx)
    assert text == ""
    assert any("Поврежденный DOCX архив" in w for w in warnings)


def test_parse_docx_empty(tmp_path: Path):
    empty_docx = tmp_path / "empty.docx"
    doc = docx.Document()
    doc.save(empty_docx)

    text, warnings = parse_document(empty_docx)
    assert text == ""
    assert any("не содержит извлекаемого текста" in w for w in warnings)


def test_parse_pdf_with_empty_page(tmp_path: Path):
    pdf_file = tmp_path / "document.pdf"

    # Создаем 2-страничный PDF: стр. 1 с текстом, стр. 2 пустая
    base_pdf = _build_minimal_pdf_bytes("FR-002 Parser Support")
    reader = PdfReader(io.BytesIO(base_pdf))
    writer = PdfWriter()
    writer.add_page(reader.pages[0])
    writer.add_blank_page(width=300, height=300)

    with pdf_file.open("wb") as f:
        writer.write(f)

    text, warnings = parse_document(pdf_file)
    assert "FR-002 Parser Support" in text
    assert any("Страницы [2]" in w for w in warnings)


def test_parse_pdf_corrupted(tmp_path: Path):
    corrupt_pdf = tmp_path / "corrupt.pdf"
    corrupt_pdf.write_bytes(b"%PDF-1.4 completely broken content")

    text, warnings = parse_document(corrupt_pdf)
    assert text == ""
    assert any("Ошибка чтения PDF" in w for w in warnings)


def test_parse_nonexistent_file(tmp_path: Path):
    missing_file = tmp_path / "not_found.md"

    text, warnings = parse_document(missing_file)
    assert text == ""
    assert any("Файл не найден" in w for w in warnings)


def test_parse_directory_as_file(tmp_path: Path):
    sub_dir = tmp_path / "sub_dir"
    sub_dir.mkdir()

    text, warnings = parse_document(sub_dir)
    assert text == ""
    assert any("не является файлом" in w for w in warnings)


def test_unsupported_format(tmp_path: Path):
    unsupported_file = tmp_path / "spec.unknown_ext"
    unsupported_file.write_text("Hello")

    with pytest.raises(ValueError, match="Неподдерживаемый формат"):
        get_parser(unsupported_file)

    text, warnings = parse_document(unsupported_file)
    assert text == ""
    assert any("Неподдерживаемый формат" in w for w in warnings)


def test_spec_parser_integration(tmp_path: Path):
    md_file = tmp_path / "integration_spec.md"
    md_file.write_text(
        "# Основные требования\n- FR-001: Реализовать парсер\n- REQ-002: Поддержать DOCX и PDF",
        encoding="utf-8"
    )

    parser = SpecParser()
    reqs, warnings = parser.parse_file(str(md_file))

    assert warnings == []
    assert len(reqs) == 2
    assert reqs[0].req_id == "FR-001"
    assert reqs[0].section == "Основные требования"
    assert reqs[1].req_id == "REQ-002"


def test_pdf_font_warnings_are_reported_without_console_flood(tmp_path, monkeypatch, caplog):
    import logging
    import specguard.rag.parsers.pdf as module
    logger = logging.getLogger('pypdf._cmap')
    original_filters = list(logger.filters)
    class Page:
        def extract_text(self):
            for _ in range(40):
                logger.warning("Got invalid hex string: %s", "Odd-length string")
            return 'FR-001: Extracted requirement'
    class Reader:
        pages = [Page()]
    monkeypatch.setattr(module, 'PdfReader', lambda path: Reader())
    path = tmp_path / 'font.pdf'
    path.touch()
    text, warnings = module.PdfParser().parse(path)
    assert 'FR-001' in text
    assert len(warnings) == 1 and '40 сообщений' in warnings[0]
    assert 'могла не извлечься' in warnings[0]
    assert 'Got invalid hex string' not in caplog.text
    assert logger.filters == original_filters
