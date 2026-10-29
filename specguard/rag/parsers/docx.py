import zipfile
from pathlib import Path

import docx
from docx.opc.exceptions import PackageNotFoundError


class DocxParser:
    """Парсер документов формата DOCX с поддержкой извлечения таблиц."""

    def parse(self, path: Path) -> tuple[str, list[str]]:
        if not path.exists():
            return "", [f"Файл не найден: {path}"]
        if not path.is_file():
            return "", [f"Указанный путь не является файлом: {path}"]

        try:
            doc = docx.Document(path)
        except (zipfile.BadZipFile, PackageNotFoundError) as exc:
            return "", [f"Поврежденный DOCX архив {path.name}: {exc}"]
        except Exception as exc:
            return "", [f"Ошибка открытия DOCX файла {path.name}: {exc}"]

        lines: list[str] = []

        for paragraph in doc.paragraphs:
            text = paragraph.text.strip()
            if text:
                lines.append(text)

        for table in doc.tables:
            for row in table.rows:
                row_cells = [
                    cell.text.replace("\r", " ").replace("\n", " ").strip()
                    for cell in row.cells
                ]
                row_text = " | ".join(row_cells).strip()
                if row_text:
                    lines.append(row_text)

        raw_text = "\n".join(lines).strip()
        warnings: list[str] = []

        if not raw_text:
            warnings.append(f"Файл {path.name} не содержит извлекаемого текста.")

        return raw_text, warnings
